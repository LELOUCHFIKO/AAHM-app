# -*- coding: utf-8 -*-
import os
import re
import json
import traceback

from kivy.app import App
from kivy.metrics import dp
from kivy.utils import get_color_from_hex
from kivy.core.window import Window
from kivy.graphics import Color, RoundedRectangle
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.screenmanager import ScreenManager, Screen
from kivy.uix.scrollview import ScrollView
from kivy.uix.gridlayout import GridLayout
from kivy.uix.filechooser import FileChooserListView

# Optional import: a missing docx dependency must not stop the app launching.
try:
    from docx import Document
    DOCX_AVAILABLE = True
    DOCX_ERROR = ""
except Exception as exc:
    Document = None
    DOCX_AVAILABLE = False
    DOCX_ERROR = str(exc)

Window.clearcolor = get_color_from_hex("#08111F")

BG = "#08111F"
PANEL = "#111F33"
PANEL2 = "#192B43"
WHITE = "#F4F7FB"
MUTED = "#9EADC1"
GOLD = "#E5B85C"
GREEN = "#2E9B72"
RED = "#D85C62"


# -------------------- DATA STORAGE --------------------

def data_path():
    """Save data in the app's private folder, not the current working directory."""
    folder = App.get_running_app().user_data_dir
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, "kursantlar.json")


def load_data():
    try:
        path = data_path()
        if not os.path.exists(path):
            return []
        with open(path, "r", encoding="utf-8") as file:
            result = json.load(file)
        return result if isinstance(result, list) else []
    except Exception:
        return []


def save_data(data):
    path = data_path()
    temp_path = path + ".tmp"
    with open(temp_path, "w", encoding="utf-8") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)
    os.replace(temp_path, path)


# -------------------- FORM FIELDS --------------------

# Display label, internal key, possible labels in a Word document.
FIELDS = [
    ("Adı", "adi", ["adı", "adi"]),
    ("Soyadı", "soyadi", ["soyadı", "soyadi"]),
    ("Ata adı", "ata_adi", ["ata adı", "ata adi"]),
    ("Doğum tarixi", "dogum_tarixi",
     ["doğulduğu tarix gün,ay,il", "doğulduğu tarix", "doğum tarixi"]),
    ("Vəsiqə seriya və FIN kodu", "fin",
     ["vəsiqə seriya və fin kodu", "vəsiqə seriya və fin", "fin kodu", "fin"]),
    ("Ümumvətəndaş pasportu", "pasport", ["ümumvətəndaş pasport"]),
    ("Doğulduğu yer", "doguldugu_yer", ["doğulduğu yer"]),
    ("Yaşayış yeri", "yasayis_yeri", ["yaşayış yeri"]),
    ("Qeydiyyat ünvanı", "qeydiyyat_unvani",
     ["qeydiyyatda olduğu ünvan"]),
    ("Qoşun növü", "qosun_novu", ["qoşun növü"]),
    ("Mülki ixtisas", "ixtisas", ["mülki ixtisas"]),
    ("İntizam balı", "intizam_bali", ["intizam balı"]),
    ("Tədris ortalama balı", "ortalama_bali", ["tədris ortalama balı"]),
    ("Fiziki hazırlığı", "fiziki_hazirliq", ["fiziki hazırlığı"]),
    ("İtirdiyi intizam balı", "itirdiyi_intizam_bali",
     ["kurs ərzində itirdiyi intizam balı"]),
    ("Kəsildiyi fənlər", "kesildiyi_fenler",
     ["son semestr imtahanı ərzində kəsildiyi fənnlər",
      "son semestr imtahanı ərzində kəsildiyi fənlər",
      "kəsildiyi fənnlər", "kəsildiyi fənlər"]),
    ("Əlaqə nömrəsi", "elaqe", ["əlaqə nömrəsi"]),
    ("Oxuduğu məktəb və ya lisey", "mekteb", ["oxuduğu məktəb və ya lisey"]),
    ("Hərbi İnstituta qəbul balı", "qebul_bali", ["hərbi instituta qəbul balı"]),
    ("Hərbi İnstituta qəbul tarixi", "qebul_tarixi", ["hərbi instituta qəbul tarixi"]),
    ("Qan qrupu", "qan_qrupu", ["qan qrupu"]),
    ("Risk qrupu", "risk_qrupu", ["risk qrupu"]),
    ("Siqaret", "siqaret", ["siqaret"]),
    ("İçki", "icki", ["içki"]),
    ("Sevimli məşğuliyyəti", "mesguliyyet", ["sevimli məşğuliyyəti"]),
    ("Ata", "ata", ["ata"]),
    ("Ata doğum tarixi", "ata_dogum_tarixi", ["ata doğum tarixi gün/ay/il", "ata doğum tarixi"]),
    ("Ata iş yeri", "ata_is_yeri", ["ata iş yeri"]),
    ("Ana", "ana", ["ana"]),
    ("Ana doğum tarixi", "ana_dogum_tarixi", ["ana doğum tarixi gün,ay,il", "ana doğum tarixi"]),
    ("Ana iş yeri", "ana_is_yeri", ["ana iş yeri"]),
    ("Bacı", "baci", ["bacı"]),
    ("Bacı doğum tarixi", "baci_dogum_tarixi", ["bacı doğum tarixi gün,ay,il", "bacı doğum tarixi"]),
    ("Bacı iş yeri", "baci_is_yeri", ["bacı iş yeri"]),
    ("Qardaş", "qardas", ["qardaş"]),
    ("Qardaş doğum tarixi", "qardas_dogum_tarixi", ["qardaş doğum tarixi gün,ay,il", "qardaş doğum tarixi"]),
    ("Qardaş iş yeri", "qardas_is_yeri", ["qardaş iş yeri"]),
    ("Ailədə vəfat edən", "vefat", ["ailədə vəfat edən"]),
    ("Boşanma halı (himayə)", "bosanma", ["boşanma halı (himayə)", "boşanma halı"]),
    ("Ailədə şəhid, qazi, əlillik", "aile_statusu",
     ["ailədə şəhid,qazi,əlillik", "ailədə şəhid, qazi, əlillik"]),
    ("Keçirdiyi xəstəliklər", "xestelikler", ["keçirdiyi xəstəliklər"]),
    ("Keçirdiyi travmalar", "travmalar", ["keçirdiyi travmalar"]),
    ("Kurs", "kurs", ["kurs"]),
    ("Tabor", "tabor", ["tabor"]),
    ("Bölük", "boluk", ["bölük", "boluk"]),
    ("Tağım", "taqim", ["tağım", "taqim"]),
    ("Qrup", "qrup", ["qrup"]),
    ("Yaxa №", "yaxa_no", ["yaxa №", "yaxa no"]),
    ("Hadisə haqqında məlumat", "hadise", ["hadisə haqqında məlumat"]),
]

DISPLAY_LABELS = {key: label for label, key, aliases in FIELDS}


# -------------------- UI HELPERS --------------------

def btn(text, bg=PANEL2, fg=WHITE, height=dp(50)):
    return Button(
        text=text,
        size_hint_y=None,
        height=height,
        background_normal="",
        background_down="",
        background_color=get_color_from_hex(bg),
        color=get_color_from_hex(fg),
        bold=True,
        font_size="14sp",
    )


class Panel(BoxLayout):
    def __init__(self, bg=PANEL, **kwargs):
        super().__init__(**kwargs)
        with self.canvas.before:
            Color(*get_color_from_hex(bg))
            self.shape = RoundedRectangle(
                pos=self.pos, size=self.size, radius=[dp(16)]
            )
        self.bind(pos=self.update_shape, size=self.update_shape)

    def update_shape(self, *_):
        self.shape.pos = self.pos
        self.shape.size = self.size


def normalise(text):
    text = str(text).lower().replace(":", "").replace("\n", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def get_aliases():
    aliases = []
    for label, key, variants in FIELDS:
        for variant in variants:
            aliases.append((normalise(variant), key))
    # Check longer labels first so "Ata doğum tarixi" isn't mistaken for "Ata".
    aliases.sort(key=lambda pair: len(pair[0]), reverse=True)
    return aliases


# -------------------- WORD IMPORT --------------------

def parse_docx(path):
    if not DOCX_AVAILABLE:
        raise RuntimeError(
            "Word dəstəyi yüklənməyib. buildozer.spec faylında "
            "requirements = python3,kivy,python-docx yazın. "
            + DOCX_ERROR[:150]
        )

    doc = Document(path)
    found = {}
    aliases = get_aliases()

    def match_label(raw_label):
        label = normalise(raw_label)
        label = re.sub(r"^\d+[\.\)]\s*", "", label)
        for alias, key in aliases:
            if label == alias:
                return key
        return None

    # Read Word tables. Supports label/value columns and "Label: value" cells.
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            for index, raw in enumerate(cells):
                if not raw:
                    continue

                label_text = raw
                value = ""

                if ":" in raw:
                    left, right = raw.split(":", 1)
                    label_text = left.strip()
                    value = right.strip()

                key = match_label(label_text)
                if not key:
                    continue

                if not value:
                    # Common form layout has a label cell followed by a value cell.
                    for next_cell in cells[index + 1:]:
                        if next_cell.strip():
                            value = next_cell.strip()
                            break

                if value:
                    found[key] = value

    # Also read paragraphs containing "Label: value".
    for paragraph in doc.paragraphs:
        if ":" not in paragraph.text:
            continue
        left, right = paragraph.text.split(":", 1)
        key = match_label(left)
        if key and right.strip():
            found[key] = right.strip()

    if not found.get("adi") and not found.get("soyadi"):
        raise ValueError(
            "Adı/Soyadı Word sənədində tapılmadı. "
            "Sənəddə məlumatların etiketləri tanınan formada olmalıdır."
        )

    found.setdefault("adi", "Naməlum")
    found.setdefault("soyadi", "Kursant")
    found.setdefault("kurs", "Hazırlıq kursu")
    found.setdefault("tabor", "Kursant hazırlıq")
    found.setdefault("boluk", "3-cü kursant")
    return found


# -------------------- LOGIN SCREEN --------------------

class Login(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)

        root = BoxLayout(
            orientation="vertical",
            padding=dp(24),
            spacing=dp(14),
        )

        root.add_widget(Label(text="", size_hint_y=0.12))

        logo = Panel(
            orientation="vertical",
            padding=dp(18),
            spacing=dp(6),
            size_hint_y=None,
            height=dp(145),
            bg="#13243A",
        )
        logo.add_widget(Label(
            text="AAHM",
            font_size="36sp",
            bold=True,
            color=get_color_from_hex(GOLD),
        ))
        logo.add_widget(Label(
            text="KURSANT ARAYIŞ-MƏRUZƏ SİSTEMİ",
            font_size="12sp",
            color=get_color_from_hex(MUTED),
        ))
        root.add_widget(logo)

        root.add_widget(Label(
            text="İdarəetmə panelinə xoş gəlmisiniz",
            color=get_color_from_hex(WHITE),
            size_hint_y=None,
            height=dp(42),
        ))

        self.username = TextInput(
            hint_text="İstifadəçi adı",
            multiline=False,
            size_hint_y=None,
            height=dp(50),
        )
        self.password = TextInput(
            hint_text="Şifrə",
            password=True,
            multiline=False,
            size_hint_y=None,
            height=dp(50),
        )

        root.add_widget(self.username)
        root.add_widget(self.password)

        self.message = Label(
            text="",
            color=get_color_from_hex(RED),
            size_hint_y=None,
            height=dp(36),
        )
        root.add_widget(self.message)

        login_button = btn("DAXİL OL  →", GOLD, "#101827", dp(54))
        login_button.bind(on_release=self.check_login)
        root.add_widget(login_button)

        root.add_widget(Label(
            text="Sınaq giriş: admin / admin",
            color=get_color_from_hex(MUTED),
            font_size="12sp",
            size_hint_y=None,
            height=dp(30),
        ))

        self.add_widget(root)

    def check_login(self, *_):
        if (
            self.username.text.strip() == "admin"
            and self.password.text == "admin"
        ):
            self.message.text = ""
            self.manager.current = "main"
        else:
            self.message.text = "İstifadəçi adı və ya şifrə yanlışdır."


# -------------------- MAIN LIST SCREEN --------------------

class Main(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)

        root = BoxLayout(
            orientation="vertical",
            padding=dp(14),
            spacing=dp(9),
        )

        top = BoxLayout(size_hint_y=None, height=dp(48))
        top.add_widget(Label(
            text="AAHM | Kursantlar",
            bold=True,
            font_size="20sp",
            color=get_color_from_hex(WHITE),
            halign="left",
        ))

        logout = btn("Çıxış", "#74343C", height=dp(40))
        logout.size_hint_x = None
        logout.width = dp(78)
        logout.bind(on_release=lambda *_: setattr(self.manager, "current", "login"))
        top.add_widget(logout)
        root.add_widget(top)

        self.search = TextInput(
            hint_text="Ad, soyad və ya ata adı ilə axtar...",
            multiline=False,
            size_hint_y=None,
            height=dp(46),
        )
        self.search.bind(text=lambda *_: self.refresh())
        root.add_widget(self.search)

        self.count = Label(
            text="KURSANTLAR: 0",
            color=get_color_from_hex(MUTED),
            size_hint_y=None,
            height=dp(25),
        )
        root.add_widget(self.count)

        scroll = ScrollView()
        self.listing = GridLayout(
            cols=1,
            spacing=dp(8),
            size_hint_y=None,
            padding=[0, dp(3)],
        )
        self.listing.bind(minimum_height=self.listing.setter("height"))
        scroll.add_widget(self.listing)
        root.add_widget(scroll)

        add_button = btn(
            "＋  WORD FAYLINDAN KURSANT ƏLAVƏ ET",
            GREEN,
            height=dp(56),
        )
        add_button.bind(
            on_release=lambda *_: setattr(self.manager, "current", "import")
        )
        root.add_widget(add_button)

        self.add_widget(root)

    def on_enter(self, *_):
        self.refresh()

    def refresh(self):
        if not hasattr(self, "listing"):
            return

        self.listing.clear_widgets()
        data = load_data()
        query = self.search.text.casefold().strip()
        count = 0

        for person in data:
            name = " ".join(
                str(person.get(key, "") or "")
                for key in ("soyadi", "adi", "ata_adi")
            ).strip()

            if query and query not in name.casefold():
                continue

            count += 1
            item_button = btn(
                "👤  " + (name or "Adsız kursant") + "   ›",
                PANEL2,
                height=dp(58),
            )
            item_button.halign = "left"
            item_button.bind(
                on_release=lambda _button, p=person: self.open_person(p)
            )
            self.listing.add_widget(item_button)

        self.count.text = "KURSANTLAR: " + str(count)

        if count == 0:
            self.listing.add_widget(Label(
                text="Siyahı boşdur. Word faylından kursant əlavə edin.",
                color=get_color_from_hex(MUTED),
                size_hint_y=None,
                height=dp(70),
            ))

    def open_person(self, person):
        detail = self.manager.get_screen("detail")
        detail.show_person(person)
        self.manager.current = "detail"


# -------------------- IMPORT SCREEN --------------------

class Import(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)

        root = BoxLayout(
            orientation="vertical",
            padding=dp(10),
            spacing=dp(8),
        )

        bar = BoxLayout(size_hint_y=None, height=dp(46))
        back = btn("‹ Geri", height=dp(42))
        back.size_hint_x = None
        back.width = dp(85)
        back.bind(on_release=lambda *_: setattr(self.manager, "current", "main"))
        bar.add_widget(back)
        bar.add_widget(Label(
            text="Word faylı seç",
            color=get_color_from_hex(WHITE),
            bold=True,
            font_size="17sp",
        ))
        root.add_widget(bar)

        root.add_widget(Label(
            text="Telefondan .docx faylını seçin, sonra ƏLAVƏ ET düyməsinə basın.",
            color=get_color_from_hex(MUTED),
            size_hint_y=None,
            height=dp(42),
            font_size="12sp",
        ))

        start_path = "/storage/emulated/0"
        if not os.path.isdir(start_path):
            start_path = App.get_running_app().user_data_dir

        self.chooser = FileChooserListView(
            path=start_path,
            filters=["*.docx", "*.DOCX"],
            dirselect=False,
            multiselect=False,
        )
        root.add_widget(self.chooser)

        self.selected_label = Label(
            text="Fayl seçilməyib",
            color=get_color_from_hex(MUTED),
            size_hint_y=None,
            height=dp(30),
            font_size="11sp",
        )
        root.add_widget(self.selected_label)
        self.chooser.bind(selection=self.selection_changed)

        import_button = btn(
            "✓  SEÇİLMİŞ WORD FAYLINI ƏLAVƏ ET",
            GREEN,
            height=dp(54),
        )
        import_button.bind(on_release=self.import_selected)
        root.add_widget(import_button)

        self.status = Label(
            text="",
            color=get_color_from_hex(WHITE),
            size_hint_y=None,
            height=dp(58),
            font_size="12sp",
        )
        root.add_widget(self.status)
        self.add_widget(root)

    def selection_changed(self, _chooser, selection):
        self.selected_label.text = (
            selection[0] if selection else "Fayl seçilməyib"
        )

    def import_selected(self, *_):
        selection = self.chooser.selection
        if not selection:
            self.status.color = get_color_from_hex(RED)
            self.status.text = "Əvvəlcə .docx faylı seçin."
            return

        try:
            person = parse_docx(selection[0])
            data = load_data()

            fin = str(person.get("fin", "")).strip().casefold()
            duplicate = False

            for old in data:
                old_fin = str(old.get("fin", "")).strip().casefold()
                same_name = (
                    str(old.get("adi", "")).casefold()
                    == str(person.get("adi", "")).casefold()
                    and str(old.get("soyadi", "")).casefold()
                    == str(person.get("soyadi", "")).casefold()
                    and str(old.get("dogum_tarixi", ""))
                    == str(person.get("dogum_tarixi", ""))
                )
                if (fin and fin == old_fin) or same_name:
                    duplicate = True
                    break

            if duplicate:
                self.status.color = get_color_from_hex(GOLD)
                self.status.text = "Bu kursant artıq siyahıdadır."
                return

            data.append(person)
            save_data(data)

            self.status.color = get_color_from_hex(GREEN)
            self.status.text = (
                "Əlavə edildi: "
                + str(person.get("soyadi", ""))
                + " "
                + str(person.get("adi", ""))
            )
            self.manager.get_screen("main").refresh()

        except Exception as exc:
            self.status.color = get_color_from_hex(RED)
            self.status.text = "Xəta: " + str(exc)[:220]


# -------------------- PERSON DETAILS AND EXPORT --------------------

class Detail(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.person = {}

        root = BoxLayout(
            orientation="vertical",
            padding=dp(12),
            spacing=dp(8),
        )

        bar = BoxLayout(size_hint_y=None, height=dp(44))
        back = btn("‹ Geri", height=dp(40))
        back.size_hint_x = None
        back.width = dp(80)
        back.bind(on_release=lambda *_: setattr(self.manager, "current", "main"))
        bar.add_widget(back)
        bar.add_widget(Label(
            text="Kursant məlumatları",
            color=get_color_from_hex(WHITE),
            bold=True,
            font_size="18sp",
        ))
        root.add_widget(bar)

        scroll = ScrollView()
        self.rows = GridLayout(
            cols=1,
            spacing=dp(5),
            size_hint_y=None,
            padding=[0, dp(3)],
        )
        self.rows.bind(minimum_height=self.rows.setter("height"))
        scroll.add_widget(self.rows)
        root.add_widget(scroll)

        export_button = btn(
            "WORD ARAYIŞ-MƏRUZƏSİ YARAT",
            GOLD,
            "#101827",
            dp(54),
        )
        export_button.bind(on_release=self.export)
        root.add_widget(export_button)

        self.status = Label(
            text="",
            color=get_color_from_hex(MUTED),
            size_hint_y=None,
            height=dp(48),
            font_size="11sp",
        )
        root.add_widget(self.status)
        self.add_widget(root)

    def show_person(self, person):
        self.person = dict(person)
        self.rows.clear_widgets()

        name = " ".join(
            str(self.person.get(key, "") or "").strip()
            for key in ("soyadi", "adi", "ata_adi")
            if str(self.person.get(key, "") or "").strip()
        )

        header = Panel(
            orientation="vertical",
            padding=dp(12),
            spacing=dp(5),
            size_hint_y=None,
            height=dp(115),
            bg="#13243A",
        )
        header.add_widget(Label(
            text=name or "Adsız kursant",
            color=get_color_from_hex(GOLD),
            font_size="19sp",
            bold=True,
            size_hint_y=None,
            height=dp(35),
        ))
        for label, key in [
            ("Kurs", "kurs"),
            ("Tabor", "tabor"),
            ("Bölük", "boluk"),
            ("Tağım", "taqim"),
        ]:
            header.add_widget(Label(
                text=label + ": " + str(self.person.get(key, "—") or "—"),
                color=get_color_from_hex(MUTED),
                font_size="12sp",
                size_hint_y=None,
                height=dp(18),
            ))
        self.rows.add_widget(header)

        hidden = {"adi", "soyadi", "ata_adi", "kurs", "tabor", "boluk", "taqim"}
        for key, value in self.person.items():
            if key in hidden:
                continue

            card = Panel(
                orientation="horizontal",
                padding=[dp(10), dp(7)],
                spacing=dp(6),
                size_hint_y=None,
                height=dp(60),
                bg=PANEL,
            )

            label_widget = Label(
                text=DISPLAY_LABELS.get(key, key),
                color=get_color_from_hex(MUTED),
                font_size="12sp",
                size_hint_x=0.46,
                halign="left",
                valign="middle",
            )
            value_widget = Label(
                text=str(value or "—"),
                color=get_color_from_hex(WHITE),
                font_size="13sp",
                size_hint_x=0.54,
                halign="left",
                valign="middle",
            )

            label_widget.bind(
                size=lambda widget, size: setattr(widget, "text_size", size)
            )
            value_widget.bind(
                size=lambda widget, size: setattr(widget, "text_size", size)
            )

            card.add_widget(label_widget)
            card.add_widget(value_widget)
            self.rows.add_widget(card)

        self.status.text = ""

    def export(self, *_):
        if not DOCX_AVAILABLE:
            self.status.color = get_color_from_hex(RED)
            self.status.text = (
                "Word dəstəyi yoxdur. buildozer.spec faylında "
                "python-docx asılılığını yoxlayın. " + DOCX_ERROR[:100]
            )
            return

        try:
            doc = Document()
            doc.add_heading("ARAYIŞ-MƏRUZƏ", 0)

            full_name = " ".join(
                str(self.person.get(key, "") or "")
                for key in ("soyadi", "adi", "ata_adi")
            ).strip()
            if full_name:
                doc.add_paragraph("Kursant: " + full_name)

            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Parametr"
            table.rows[0].cells[1].text = "Məlumat"

            for key, value in self.person.items():
                cells = table.add_row().cells
                cells[0].text = DISPLAY_LABELS.get(key, key)
                cells[1].text = str(value or "")

            folder = os.path.join(
                App.get_running_app().user_data_dir,
                "Arayislar",
            )
            os.makedirs(folder, exist_ok=True)

            surname = str(self.person.get("soyadi", "Kursant"))
            safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", surname).strip("_")
            if not safe_name:
                safe_name = "Kursant"

            path = os.path.join(folder, "Arayis_" + safe_name + ".docx")
            doc.save(path)

            self.status.color = get_color_from_hex(GREEN)
            self.status.text = (
                "Word faylı yaradıldı. Tətbiqin şəxsi Arayislar qovluğunda saxlanıldı."
            )

        except Exception as exc:
            self.status.color = get_color_from_hex(RED)
            self.status.text = "Word yaratma xətası: " + str(exc)[:180]


# -------------------- APP --------------------

class AAHMApp(App):
    def build(self):
        self.title = "AAHM Kursant Sistemi"

        manager = ScreenManager()
        manager.add_widget(Login(name="login"))
        manager.add_widget(Main(name="main"))
        manager.add_widget(Import(name="import"))
        manager.add_widget(Detail(name="detail"))

        return manager


if __name__ == "__main__":
    try:
        AAHMApp().run()
    except Exception:
        # Print startup traceback into the Android/build log for diagnosis.
        traceback.print_exc()
